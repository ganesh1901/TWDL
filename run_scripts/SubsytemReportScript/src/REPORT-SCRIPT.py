#!/usr/bin/python


from matplotlib.pyplot import figure, show

from numpy import arange, sin, pi


import openpyxl
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.chart import LineChart,Reference, Series
from openpyxl.chart.shapes import GraphicalProperties

import struct
import os
import time
import sys

from ReportVarList import *

timetag_fmt = 'L'
header_fmt = 'H'


if(len(sys.argv) != 2):
    print 'usage == python 1)Command Name '
    exit(0)

cmd_name = sys.argv[1]
try:
        input_file_list=logfile_map[cmd_name]
        for i in range(len(input_file_list)):
                real_path = os.readlink(input_file_list[i])
               
                time1 = os.path.getctime(real_path)
                timestamp = time.strftime('%d-%m-%Y-%H-%M-%S', time.localtime(time1))
               
                index = index_map[cmd_name]
		


                wb = load_workbook(os.getenv('PWD')+'/../inp/template.xlsx')
                ws = wb[sheet_name_map[index[i]]]

                fd_out = open(sheet_name_map[index[i]]+str('.csv'), "w+")


                timetag_size = struct.calcsize(timetag_fmt)
                header_size = struct.calcsize(header_fmt)
                packet_size = timetag_size + header_size + 64

      
                line_count = os.stat(input_file_list[i]).st_size / packet_size
		print "\n*************Report Generation Start*******************"
                print " For ",sheet_name_map[index[i]],' with line count ', line_count
                
                f = open(input_file_list[i], "rb")
               
                for y in range(line_count):
                        bytes = f.read(packet_size)
                        if bytes:
                               
                                data = struct.unpack(fmt_list[index[i]], bytes[timetag_size + header_size:  74])
                                #print data
                                ti=[]
                                try:
                                        for k in range(len(index_list[index[i]])):
                                                if(scale_list[index[i]][k]==-1):
                                                        ti.append(data[index_list[index[i]][k]])
                                                else:
                                                        ti.append(data[index_list[index[i]][k]] * scale_list[index[i]][k])
                                                #ti.append('\t')

                                        #ti.append('\n')

                                        for jj in ti:
                                                fd_out.write(str(jj))


                                        ws.append(ti)
                                        #print "line number=",y

                                except IndexError as I:
                                        print "Error  ",i

                        else:
                            #print ' Ended'
                            break
                            
		for key, val in sheet_name_map.items():
			if val!=sheet_name_map[index[i]]:
				sheet_ext=wb[val]#.get_sheet_by_name(val)
				wb.remove(sheet_ext)
                wb.save(input_file_list[i]+"-report" + '.xlsx')
                fd_out.close()

                print "*************Report Generation End*********************\n\n"
	

except KeyError as K:
        print "KeyError:",K
        
except OSError as E:
        print "OSError: ",E


